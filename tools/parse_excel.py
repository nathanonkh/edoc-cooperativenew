# -*- coding: utf-8 -*-
"""อ่านไฟล์ Excel ครุภัณฑ์ 3 ไฟล์ -> รายการ dict (ใช้สร้าง seed_assets.sql)"""
import re, glob, datetime
import openpyxl, xlrd

TH_MON = {'มค':1,'กพ':2,'มีค':3,'เมย':4,'พค':5,'มิย':6,'กค':7,'สค':8,'กย':9,'ตค':10,'พย':11,'ธค':12}

def be_from_year(y):
    """ปีที่ Excel เก็บมา -> ปี พ.ศ. (ไฟล์เดิมปนกัน: 1950-1999 = พ.ศ. 2 หลักที่ Excel อ่านผิดเป็น ค.ศ., 2400+ = พ.ศ., 2000-2399 = ค.ศ.)"""
    if 1950 <= y <= 1999: return 2500 + (y - 1900)
    if y >= 2400: return y
    return y + 543

def thai_date(v, flags, ctx):
    """คืนสตริง dd/mm/yyyy (พ.ศ.) หรือ '' """
    if v is None or v == '': return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return '%02d/%02d/%d' % (v.day, v.month, be_from_year(v.year))
    if isinstance(v, (int, float)):
        d = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(v))
        return '%02d/%02d/%d' % (d.day, d.month, be_from_year(d.year))
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})\s*([ก-๙\.\s]+?)\s*(\d{2,4})$', s)
    if m:
        mon = TH_MON.get(re.sub(r'[\.\s]', '', m.group(2)))
        y = int(m.group(3)); y = y + 2500 if y < 100 else y
        if mon: return '%02d/%02d/%d' % (int(m.group(1)), mon, y)
    flags.append('%s: วันที่อ่านไม่ได้ "%s"' % (ctx, s))
    return ''

def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(',', '').strip()
    try: return float(s)
    except: return None

COND = {'ปกติ':'normal','ชำรุด':'damaged','เสื่อมสภาพ':'deteriorated','ทรุดโทรม':'dilapidated','สูญหาย':'lost','จำหน่ายแล้ว':'disposed'}
def cond(v, flags, ctx):
    s = ('' if v is None else str(v)).strip()
    if s in COND: return COND[s]
    if s: flags.append('%s: สภาพ "%s" ไม่รู้จัก ใช้ปกติ' % (ctx, s))
    return 'normal'

def s(v): return '' if v is None else re.sub(r'\s+', ' ', str(v)).strip()

CAT_FIX = {'คุรภัณฑ์ยานพาหนะ': 'ครุภัณฑ์ยานพาหนะ'}

def rows_from(sheet_rows, first, is_item, make):
    cat = ''; out = []
    for i, r in enumerate(sheet_rows):
        if i < first: continue
        r = list(r) + [None] * 14
        if is_item(r): out.append(make(r, cat, i + 1))
        else:
            label = s(r[1]) if s(r[1]) and not s(r[0]) and not any(s(x) for x in r[2:10]) else ''
            label = label or (s(r[2]) if s(r[2]) and not s(r[0]) and not any(s(x) for x in r[3:10]) else '')
            if label and 'ข้าพเจ้า' not in label and 'ลงชื่อ' not in label and not label.startswith('('):
                cat = CAT_FIX.get(label, label)
    return out

def load(upload_dir='/mnt/user-data/uploads'):
    flags = []; data = []
    # 1) ครุภัณฑ์ (.xls) ชีต ส่งกรม68
    f1 = glob.glob(upload_dir + '/1_*.xls')[0]
    sh = xlrd.open_workbook(f1).sheet_by_name('ส่งกรม68')
    rows = []
    for i in range(sh.nrows):
        r = []
        for c in range(sh.ncols):
            cell = sh.cell(i, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                r.append(xlrd.xldate_as_datetime(cell.value, 0))
            else:
                r.append(cell.value if cell.ctype != xlrd.XL_CELL_EMPTY else None)
        rows.append(r)
    def item1(r): return isinstance(r[0], (int, float)) or (s(r[0]).isdigit()) or (s(r[0]) == '' and s(r[2]) and s(r[3]) and s(r[1]))
    def mk1(r, cat, line):
        ctx = 'ไฟล์1 แถว%d' % line
        if s(r[2]) and r[2] not in (None, ''): gf = s(int(r[2])) if isinstance(r[2], float) else s(r[2])
        else: gf = ''
        return dict(asset_type='registered', category=cat, name=s(r[1]), gf_code=gf, official_no=s(r[3]),
                    quantity=1, unit='', price=num(r[5]) or 0, acquired_date=thai_date(r[4], flags, ctx),
                    acquire_method=s(r[6]), acq_cost=num(r[7]), depreciation=num(r[8]),
                    status=cond(r[9], flags, ctx), note=s(r[10]), src=ctx)
    data += rows_from(rows, 4, item1, mk1)

    # 2) ครุภัณฑ์ไม่เกิน 10,000 ชีต ส่งกรม69  (คอลัมน์ L = จำหน่ายแล้ว)
    wb = openpyxl.load_workbook(glob.glob(upload_dir + '/2_*.xlsx')[0], data_only=True)
    ws = [list(r) for r in wb['ส่งกรม69'].iter_rows(values_only=True)]
    def item2(r): return isinstance(r[0], (int, float))
    def mk2(r, cat, line):
        ctx = 'ไฟล์2 แถว%d' % line
        st = cond(r[9], flags, ctx)
        if s(r[11]): st = 'disposed'
        return dict(asset_type='under10k', category=cat or s(r[1]), name=s(r[2]), gf_code='', official_no=s(r[3]),
                    quantity=int(num(r[4]) or 1), unit=s(r[5]), price=num(r[6]) or 0,
                    acquired_date=thai_date(r[7], flags, ctx), acquire_method=s(r[8]),
                    acq_cost=None, depreciation=None, status=st, note=s(r[10]) if len(r) > 10 else '', src=ctx,
                    disposed_mark=s(r[11]))
    data += rows_from(ws, 4, item2, mk2)

    # 3) ครุภัณฑ์นอกทะเบียน ชีต ตรวจนับ69
    wb = openpyxl.load_workbook(glob.glob(upload_dir + '/3_*.xlsx')[0], data_only=True)
    ws = [list(r) for r in wb['ตรวจนับ69'].iter_rows(values_only=True)]
    def item3(r): return isinstance(r[0], (int, float))
    def mk3(r, cat, line):
        ctx = 'ไฟล์3 แถว%d' % line
        return dict(asset_type='external', category=cat, name=s(r[1]), gf_code='', official_no=s(r[2]),
                    quantity=int(num(r[3]) or 1), unit=s(r[4]), price=0, acquired_date='', acquire_method='',
                    acq_cost=None, depreciation=None, status=cond(r[5], flags, ctx), note=s(r[6]), src=ctx)
    data += rows_from(ws, 4, item3, mk3)
    return data, flags

if __name__ == '__main__':
    d, fl = load()
    from collections import Counter
    print(len(d), Counter(x['asset_type'] for x in d))
    print(Counter((x['asset_type'], x['category']) for x in d))
    print(Counter((x['asset_type'], x['status']) for x in d))
    print('flags', len(fl)); [print(' ', f) for f in fl[:40]]
    for t in ('registered','under10k','external'):
        print([x for x in d if x['asset_type']==t][:2])
